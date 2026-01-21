from datetime import datetime, date
from .models import *
import csv, io, re
from openpyxl import Workbook
from io import BytesIO

def to_midnight(d):
    return datetime(d.year, d.month, d.day)

def zip_object_factory(npci_username):
    try:
        seq = Zip.objects.filter(date__gte=date.today()).latest("seq_no").seq_no + 1
    except Zip.DoesNotExist:
        seq = 1

    filename = 'MMS-CREATE-HGBX-' + npci_username + '-' + date.today().strftime(r'%d%m%Y') + '-' + str(seq).zfill(6) + '-INP.zip'

    zip_obj = Zip(
        date = date.today(),
        seq_no = seq,
        npci_username = npci_username,
        filename = filename
    )

    zip_obj.save()

    return zip_obj

def presentation_object_factory(npci_username):
    try:
        seq = Presentation.objects.filter(date__gte=date.today()).latest("seq_no").seq_no + 1
    except Presentation.DoesNotExist:
        seq = 1

    filename_prefix = 'MMS-CREATE-HGBX-' + npci_username + '-' + date.today().strftime(r'%d%m%Y') + '-' + str(seq).zfill(6)

    npci_MsgId = 'HGBX' + date.today().strftime(r'%d%m%Y') + str(seq).zfill(6)

    presentation_obj = Presentation(
        date = date.today(),
        seq_no = seq,
        npci_username = npci_username,
        npci_MsgId = npci_MsgId,
        filename_prefix = filename_prefix
    )

    return presentation_obj

def get_presentation_from_filenames(xml_dict, zip_file_name):
    zip_file_date = datetime.strptime(zip_file_name.split("-")[4], r"%d%m%Y").date()
    zip_file_seq = int(zip_file_name.split("-")[5])

    xml_file_name = xml_dict['filename']
    xml_file_date = datetime.strptime(xml_file_name.split("-")[4], r"%d%m%Y").date()
    xml_file_seq = int(xml_file_name.split("-")[5])

    zip_obj = Zip.objects.get(date = zip_file_date, seq_no = zip_file_seq)
    presentation_obj = Presentation.objects.get(zip = zip_obj, date = xml_file_date, seq_no = xml_file_seq)

    return presentation_obj


def process_ack(file, zip_filename):
    status_dict = {
        'filename': None,
        'found': False,
        'status': None
    }
    status_dict['filename'] = file['filename']
    try:
        p = Presentation.objects.get(npci_MsgId = file['OriginalMsgId'])
    except KeyError:
        try:
            p = get_presentation_from_filenames(file, zip_filename)
        except (Presentation.DoesNotExist, Zip.DoesNotExist):
            return status_dict
    except Presentation.DoesNotExist:
        return status_dict
    
    status_dict['found'] = True

    if p.npci_upload_time:
        print('already updated')
        status_dict['status'] = "Status Already Updated"
        return status_dict
    
    p.npci_upload_time = file['AcqCreDtTm']
    
    p.mandate.init_req_flag = False
    p.mandate.save()
    print('Init req flag updated, mandate saved.')

    if file['Accptd'] == 'true':
        p.npci_umrn = file['UMRN']
        status_dict['status'] = 'UMRN: ' + file['UMRN']
    elif file['Accptd'] == 'false':
        p.npci_upload_error = file['Error']
        status_dict['status'] = 'Error: ' + file['Error']
    else:
        status_dict['status'] = 'Error.'
    p.save()
    
    return status_dict


def process_status(file):
    filerow = io.StringIO(file.read().decode('utf-8'))
    dictreader = csv.DictReader(filerow)
    messages = []
    for row in dictreader:
        require_save = False
        message = {
            'umrn': None,
            'status': None,
            'code': None,
            'save': False
        }
        try:
            p = Presentation.objects.get(npci_umrn = row['UMRN'])
            message['umrn'] = row['UMRN']
            
            #updating status (active/rejected)
            if p.npci_status == None:
                p.npci_status = row['Status']
                require_save = True
                message['status'] = "New status: " + p.npci_status
            else:
                # status already updated
                message['status'] = "Status already updated."
            
            #updating code
            try:
                p.set_reason_code(row['Code'])
                require_save = True
                message['code'] = "New response code: " + p.npci_reason_code
            except KeyError:
                message['code'] = "'Code' not found in response file."
            except ValueError as err:
                message['code'] = str(err)
            
            if require_save:
                p.save()
                message['save'] = True

        except Presentation.DoesNotExist:
            message['status'] = "UMRN not found in Presentation table"
        except KeyError:
            message['umrn'] = 'UMRN key not found in the dict'
            break
        
        messages.append(message)
    
    return messages


# Getting the OFFICE queryset based on the user
def get_office_queryset(office):
    queryset = Office.objects.filter(type='BO')
    if office.type == 'RO':
        return queryset.filter(region = office.region)
    elif office.type == 'BO':
        return queryset.filter(sol_id = office.sol_id)
    else:
        return queryset
    
# Getting the MANDATE queryset based on the user
def get_mandate_queryset(user_office):
    queryset = Mandate.objects.filter(is_deleted = False)
    if user_office.type == 'HO':
        return queryset
    if user_office.type == 'RO':
        return queryset.filter(office__region = user_office.region)
    elif user_office.type == 'BO':
        return queryset.filter(office = user_office)
    

#Checking if user has permission over the mandate
def user_mandate_allowed(user, mandate):
    user_office = user.userextended.office

    if user_office.type == 'HO':
        return True
    if user_office.type == 'RO' and mandate.office.region == user_office.region:
        return True
    elif user_office.type == 'BO' and user_office == mandate.office:
        return True
    
    return False

# for creating and returning the excel file bytes for a particular date
def create_excel_debit_list(dt):

    workbook = Workbook()
    worksheet = workbook.active
    # worksheet = workbook.create_sheet()

    header = [
        'UMRN',
        'Date of Manate',
        'Start Date',
        'End Date',
        'Debtor Name',
        'Debtor Bank',
        'Debtor Account',
        'Debtor IFSC',
        'Credit Account',
        'Amount',
        'Ref no.',
        'Cancel Request Flag',
        'Cancel Request User',
        'Cancel Request Time'
    ]
    worksheet.append(header)

    for p in Presentation.objects.filter(npci_status='Active', cancel_flg=False).filter(mandate__debit_date=dt):
    
        list = (
            p.npci_umrn,
            p.mandate.date,
            p.mandate.start_date,
            p.mandate.end_date,
            p.mandate.debtor_name,
            p.mandate.debtor_bank.name,
            p.mandate.debtor_acc_no,
            p.mandate.debtor_acc_ifsc,
            p.mandate.credit_account,
            p.mandate.amount,
            p.mandate.ref,
            p.cancel_req_flg,
            p.cancel_req_user.username if p.cancel_req_user != None else None,
            p.cancel_req_time.replace(tzinfo=None) if p.cancel_req_time != None else None
        )
        
        # append this row to sheet
        worksheet.append(list)
    
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    buffer.seek(0)
    return buffer